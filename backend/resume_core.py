"""
Core resume parsing + scoring logic.
Kept separate from app.py so it can be tested or reused on its own.
"""

import os
import re
import pdfplumber
import docx
import spacy

from skills_list import SKILLS_DB, ROLE_SKILLS

try:
    nlp = spacy.load("en_core_web_sm")
except OSError:
    raise RuntimeError(
        "spaCy model missing. Run: python -m spacy download en_core_web_sm"
    )

EMAIL_REGEX = r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}"
PHONE_REGEX = r"(\+?\d[\d\s\-()]{8,14}\d)"
EDUCATION_KEYWORDS = [
    "b.tech", "btech", "b.e", "m.tech", "mtech", "bachelor", "master",
    "phd", "ph.d", "mba", "bca", "mca", "diploma", "high school",
    "intermediate", "b.sc", "m.sc", "university", "college", "institute",
]
EXPERIENCE_REGEX = r"(\d+)\+?\s+years?\s+(of\s+)?experience"


def extract_text(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".pdf":
        text = ""
        with pdfplumber.open(filepath) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
                for h in (page.hyperlinks or []):
                    if h.get("uri"):
                        text += "\n" + h["uri"]
        return text
    elif ext in (".docx", ".doc"):
        d = docx.Document(filepath)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        for rel in d.part.rels.values():
            if "hyperlink" in rel.reltype:
                parts.append(rel.target_ref)
        return "\n".join(parts)
    else:
        raise ValueError(f"Unsupported file type: {ext}")


def extract_name(text, doc):
    # Resumes almost always start with the name: use the first line if it looks like one
    for line in text.split("\n")[:3]:
        line = line.strip()
        parts = line.split()
        if 2 <= len(parts) <= 4 and all(re.fullmatch(r"[A-Za-z.'-]+", w) for w in parts) \
                and not any(k in line.lower() for k in ("resume", "curriculum", "education", "skills", "b.tech")):
            return line.title() if line.isupper() else line
    first_chunk = "\n".join(text.split("\n")[:10])
    chunk_doc = nlp(first_chunk)
    for ent in chunk_doc.ents:
        if ent.label_ == "PERSON":
            return ent.text.strip()
    for ent in doc.ents:
        if ent.label_ == "PERSON":
            return ent.text.strip()
    return None


def extract_email(text):
    match = re.search(EMAIL_REGEX, text)
    return match.group(0) if match else None


def extract_phone(text):
    candidates = re.findall(PHONE_REGEX, text)
    for c in candidates:
        digits = re.sub(r"\D", "", c)
        if 10 <= len(digits) <= 13:
            return c.strip()
    return None


ALL_SKILLS = sorted(set(SKILLS_DB) | {x for v in ROLE_SKILLS.values() for x in v})


def extract_skills(text, extra=None):
    text_lower = text.lower()
    found = set()
    for skill in set(ALL_SKILLS) | set(extra or []):
        pattern = r"(?<![a-zA-Z])" + re.escape(skill) + r"(?![a-zA-Z])"
        if re.search(pattern, text_lower):
            found.add(skill)
    return sorted(found)


def extract_education(text):
    found = []
    for line in text.split("\n"):
        line_lower = line.lower()
        if any(kw in line_lower for kw in EDUCATION_KEYWORDS):
            cleaned = line.strip()
            if cleaned:
                found.append(cleaned)
    return found[:5]


def extract_experience(text):
    match = re.search(EXPERIENCE_REGEX, text.lower())
    if match:
        return f"{match.group(1)}+ years"
    return None


def score_resume(parsed):
    """ATS-style completeness score out of 100."""
    score = 0
    breakdown = {}

    breakdown["contact_info"] = 15 if (parsed["email"] and parsed["phone"]) else (
        8 if (parsed["email"] or parsed["phone"]) else 0
    )
    breakdown["name_detected"] = 10 if parsed["name"] else 0
    breakdown["education"] = 20 if parsed["education"] else 0
    breakdown["experience"] = 15 if parsed["experience"] else 0

    skill_count = len(parsed["skills"])
    breakdown["skills"] = min(40, skill_count * 4)  # up to 10 skills = full marks

    score = sum(breakdown.values())
    return {"total": min(score, 100), "breakdown": breakdown}


def role_fit(parsed, role, custom=None):
    """Percentage match against a role's skill set (or custom skills)."""
    if custom:
        required, role = set(custom), "your custom skills"
    elif role in ROLE_SKILLS:
        required = set(ROLE_SKILLS[role])
    else:
        return None

    have = set(parsed["skills"])
    matched = required & have
    missing = required - have

    pct = round(100 * len(matched) / len(required)) if required else 0
    return {
        "role": role,
        "match_percent": pct,
        "matched_skills": sorted(matched),
        "missing_skills": sorted(missing),
    }


ACTION_VERBS = {"developed","built","designed","implemented","created","led","managed",
 "optimized","analyzed","deployed","engineered","improved","launched","automated",
 "integrated","trained","collaborated","delivered","reduced","increased"}
SECTION_KEYS = {
    "Summary": ["summary", "objective", "profile"],
    "Education": ["education"],
    "Experience": ["experience", "internship", "work history"],
    "Projects": ["projects", "project"],
    "Skills": ["skills"],
    "Certifications": ["certification", "certificate", "courses"],
    "Achievements": ["achievement", "awards", "honors"],
}


def compute_insights(text):
    lines = [l.strip().lower() for l in text.split("\n") if l.strip()]
    heads = [l for l in lines if len(l) < 40]
    sections = {k: any(any(w in h for w in ws) for h in heads) for k, ws in SECTION_KEYS.items()}
    low = text.lower()
    words = re.findall(r"[a-zA-Z']+", low)
    return {
        "word_count": len(words),
        "action_verbs": len({w for w in words if w in ACTION_VERBS}),
        "metrics": len(re.findall(r"\d+(?:\.\d+)?\s*(?:%|\+|k\b|x\b)", low)),
        "sections": sections,
        "links": {"GitHub": "github" in low, "LinkedIn": "linkedin" in low},
    }


def generate_suggestions(parsed, score, rf=None, ins=None):
    """Plain-language, actionable fixes based on what's missing."""
    tips = []

    if not parsed["email"]:
        tips.append("Add a clearly formatted email address near the top of your resume.")
    if not parsed["phone"]:
        tips.append("Add a phone number near the top — ATS systems look for this in the header.")
    if not parsed["name"]:
        tips.append("Make sure your full name appears as plain text at the very top (not inside an image or text box — those aren't machine-readable).")
    if not parsed["education"]:
        tips.append("Add a clearly labeled Education section with your degree, institution, and year.")
    if not parsed["experience"]:
        tips.append("State your experience explicitly, e.g. '2+ years of experience in web development', so it's easy to detect.")

    skill_count = len(parsed["skills"])
    if skill_count == 0:
        tips.append("Add a dedicated Skills section listing your technical skills — none were detected right now.")
    elif skill_count < 6:
        tips.append(f"Only {skill_count} recognizable skills were found. List more specific tools/technologies (e.g. exact language, framework, or library names) rather than vague phrases.")

    if rf and rf["missing_skills"]:
        top_missing = ", ".join(rf["missing_skills"][:5])
        tips.append(f"For '{rf['role']}', consider adding or highlighting: {top_missing}.")

    if ins:
        if not ins["sections"]["Projects"]:
            tips.append("Add a Projects section with 2-3 projects, the tech used, and what each achieved.")
        if not ins["links"]["GitHub"]:
            tips.append("Add your GitHub profile link so recruiters can see your code.")
        if not ins["links"]["LinkedIn"]:
            tips.append("Add your LinkedIn URL to the header.")
        if ins["metrics"] == 0:
            tips.append("Quantify your impact with numbers, e.g. 'reduced load time by 40%' or 'served 500+ users'.")
        if ins["action_verbs"] < 3:
            tips.append("Start bullet points with strong action verbs like Built, Designed, Deployed, Optimized.")
        if ins["word_count"] < 150:
            tips.append("Your resume looks thin. Aim for 250-600 words with specific details.")
        elif ins["word_count"] > 900:
            tips.append("Your resume is long. Trim it to 1-2 pages, keeping only the most relevant points.")

    if score["total"] >= 85 and not tips:
        tips.append("Strong resume — all key sections and a solid skill count were detected.")

    return tips


def analyze(filepath, target_role=None, custom_skills=None):
    text = extract_text(filepath)
    doc = nlp(text)

    parsed = {
        "name": extract_name(text, doc),
        "email": extract_email(text),
        "phone": extract_phone(text),
        "skills": extract_skills(text, custom_skills),
        "education": extract_education(text),
        "experience": extract_experience(text),
    }

    ins = compute_insights(text)
    score = score_resume(parsed)
    rf = role_fit(parsed, target_role, custom_skills) if (target_role or custom_skills) else None

    result = {
        "parsed": parsed,
        "score": score,
        "available_roles": list(ROLE_SKILLS.keys()),
        "insights": ins,
        "suggestions": generate_suggestions(parsed, score, rf, ins),
    }

    if rf:
        result["role_fit"] = rf

    return result
